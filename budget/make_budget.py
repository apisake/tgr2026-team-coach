import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = sys.argv[1]
FONT = "Tahoma"
BLUE = "0000FF"
GREEN = "008000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="D9E1F2")
TOTAL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="A6A6A6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BAHT = '#,##0;(#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'


def f(bold=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


def style_range(ws, rng, font=None, fill=None, fmt=None, border=True, align=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if fmt: c.number_format = fmt
            if border: c.border = BOX
            if align: c.alignment = align


wb = Workbook()

# ---------------- Sheet 1: assumptions ----------------
a = wb.active
a.title = "สมมติฐาน"
a["A1"] = "สมมติฐานงบประมาณค่าอาหาร ทีม TGR2026"
a["A1"].font = f(True, size=13)
rows = [
    ("จำนวนนักศึกษา", 15, "คน"),
    ("จำนวนอาจารย์", 3, "คน"),
    ("จำนวนคนรวม", "=B3+B4", "คน"),
    ("อาหารว่าง: จำนวนมื้อต่อวัน", 2, "มื้อ/วัน"),
    ("อาหารว่าง: ราคาต่อคนต่อมื้อ", 50, "บาท"),
    ("อาหารกลางวัน: จำนวนมื้อต่อวัน", 1, "มื้อ/วัน"),
    ("อาหารกลางวัน: ราคาต่อคนต่อมื้อ", 100, "บาท"),
    ("อาหารเย็น: ราคาต่อคนต่อมื้อ (เฉพาะ 24 ต.ค. 2569)", 100, "บาท"),
    ("จำนวนวันจัดเลี้ยง", 4, "วัน"),
    ("ราคาจริง: อาหารว่างต่อคนต่อมื้อ", 40, "บาท"),
    ("ราคาจริง: อาหารกลางวันต่อคนต่อมื้อ", 55, "บาท"),
    ("ราคาจริง: อาหารเย็นต่อคนต่อมื้อ", 55, "บาท"),
]
a["A2"], a["B2"], a["C2"] = "รายการ", "ค่า", "หน่วย"
style_range(a, "A2:C2", font=f(True), fill=HEAD)
for i, (lab, val, unit) in enumerate(rows, start=3):
    a[f"A{i}"], a[f"B{i}"], a[f"C{i}"] = lab, val, unit
    is_input = not (isinstance(val, str) and val.startswith("="))
    style_range(a, f"A{i}:C{i}", font=f())
    a[f"B{i}"].font = f(color=BLUE if is_input else "000000")
    a[f"B{i}"].number_format = NUM
    if is_input:
        a[f"B{i}"].fill = YELLOW
a["B3"].comment = Comment("ตัวเลขจากผู้ใช้: นักศึกษา 15 คน", "budget")
a["B4"].comment = Comment("ตัวเลขจากผู้ใช้: อาจารย์ 3 คน", "budget")
a["B12"].comment = Comment("ราคาที่ทำได้จริง จากผู้ใช้", "budget")

a["A17"] = "วิธีใช้"
a["A17"].font = f(True)
notes = [
    "ตัวเลขทั้งหมดในชีตนี้มาจากโจทย์ที่ผู้ใช้กำหนด",
    "ช่องสีเหลืองตัวอักษรสีน้ำเงิน = ค่าที่แก้ได้ ทุกชีตจะคำนวณใหม่อัตโนมัติ",
    "วันจัดเลี้ยง: 23, 24, 25 ต.ค. 2569 และ 1 พ.ย. 2569 (อาหารว่าง 2 มื้อ + กลางวัน ทุกวัน)",
    "อาหารเย็นมีเพียงวันเดียว คือ 24 ต.ค. 2569",
    "ราคาจริง (แถว 12-14) ใช้คำนวณยอดจ่ายจริง แยกจากราคาตั้งงบ",
    "ชีต 'งบประมาณ': กรอกจำนวนคนที่กินจริงในช่องสีเหลือง เพื่อดูยอดจ่ายจริงและส่วนต่าง",
    "ชีต 'จำนวนเงินที่เป็นไปได้': ยอดเงินทุกกรณีตามจำนวนคน 0-18 คน",
]
for i, n in enumerate(notes, start=18):
    a[f"A{i}"] = "• " + n
    a[f"A{i}"].font = f()
a.column_dimensions["A"].width = 48
a.column_dimensions["B"].width = 12
a.column_dimensions["C"].width = 10

P = "สมมติฐาน!"
PEOPLE = f"{P}$B$5"
SNACK = f"{P}$B$7"
LUNCH = f"{P}$B$9"
DINNER = f"{P}$B$10"
SNACK_A = f"{P}$B$12"
LUNCH_A = f"{P}$B$13"
DINNER_A = f"{P}$B$14"

# ---------------- Sheet 2: budget vs actual ----------------
b = wb.create_sheet("งบประมาณ")
b["A1"] = "งบประมาณค่าอาหาร เทียบกับยอดจ่ายจริง"
b["A1"].font = f(True, size=13)
heads = ["วันที่", "มื้อ", "ราคาตั้งงบ/คน (บาท)", "จำนวนคน (ตั้งงบ)", "งบประมาณ (บาท)",
         "ราคาจริง/คน (บาท)", "จำนวนคน (จริง)", "จ่ายจริง (บาท)", "ส่วนต่าง (บาท)\nงบ − จริง"]
for j, h in enumerate(heads):
    b.cell(row=2, column=j + 1, value=h)
style_range(b, "A2:I2", font=f(True), fill=HEAD,
            align=Alignment(horizontal="center", vertical="center", wrap_text=True))
b.row_dimensions[2].height = 32

days = ["23 ต.ค. 2569", "24 ต.ค. 2569", "25 ต.ค. 2569", "1 พ.ย. 2569"]
r = 3
for d in days:
    meals = [("อาหารว่างเช้า", SNACK, SNACK_A), ("อาหารกลางวัน", LUNCH, LUNCH_A),
             ("อาหารว่างบ่าย", SNACK, SNACK_A)]
    if d.startswith("24"):
        meals.append(("อาหารเย็น", DINNER, DINNER_A))
    for m, price, price_a in meals:
        b[f"A{r}"], b[f"B{r}"] = d, m
        b[f"C{r}"] = f"={price}"
        b[f"D{r}"] = f"={PEOPLE}"
        b[f"E{r}"] = f"=C{r}*D{r}"
        b[f"F{r}"] = f"={price_a}"
        b[f"G{r}"] = f"=D{r}"
        b[f"H{r}"] = f"=F{r}*G{r}"
        b[f"I{r}"] = f"=E{r}-H{r}"
        style_range(b, f"A{r}:I{r}", font=f(), fmt=BAHT)
        for col in "CDF":
            b[f"{col}{r}"].font = f(color=GREEN)
        b[f"G{r}"].fill = YELLOW
        r += 1
last = r - 1
b[f"A{r}"] = "รวมทั้งหมด"
b.merge_cells(f"A{r}:C{r}")
for col in "DEGHI":
    b[f"{col}{r}"] = f"=SUM({col}3:{col}{last})"
style_range(b, f"A{r}:I{r}", font=f(True), fill=TOTAL, fmt=BAHT)
b["G3"].comment = Comment("กรอกจำนวนคนที่กินจริงของแต่ละมื้อ ค่าเริ่มต้นเท่ากับจำนวนที่ตั้งงบ", "budget")
tot = r
r += 1
b[f"A{r}"] = "ส่วนต่างคิดเป็น % ของงบ"
b.merge_cells(f"A{r}:H{r}")
b[f"I{r}"] = f"=IF(E{tot}=0,0,I{tot}/E{tot})"
style_range(b, f"A{r}:I{r}", font=f(True), fill=TOTAL, fmt='0.0%;(0.0%);"-"')

r += 2
b[f"A{r}"] = "สรุปตามประเภทอาหาร"
b[f"A{r}"].font = f(True)
r += 1
for j, h in enumerate(["ประเภท", "จำนวนมื้อ (คน-มื้อ)", "งบประมาณ (บาท)", "จ่ายจริง (บาท)", "ส่วนต่าง (บาท)"]):
    b.cell(row=r, column=j + 1, value=h)
style_range(b, f"A{r}:E{r}", font=f(True), fill=HEAD, align=Alignment(horizontal="center", wrap_text=True))
r += 1
s0 = r
for label, crit in [("อาหารว่าง", "อาหารว่าง*"), ("อาหารกลางวัน", "อาหารกลางวัน"), ("อาหารเย็น", "อาหารเย็น")]:
    rng = f"$B$3:$B${last}"
    b[f"A{r}"] = label
    b[f"B{r}"] = f'=SUMIF({rng},"{crit}",$D$3:$D${last})'
    b[f"C{r}"] = f'=SUMIF({rng},"{crit}",$E$3:$E${last})'
    b[f"D{r}"] = f'=SUMIF({rng},"{crit}",$H$3:$H${last})'
    b[f"E{r}"] = f"=C{r}-D{r}"
    style_range(b, f"A{r}:E{r}", font=f(), fmt=BAHT)
    r += 1
b[f"A{r}"] = "รวม"
for col in "BCDE":
    b[f"{col}{r}"] = f"=SUM({col}{s0}:{col}{r-1})"
style_range(b, f"A{r}:E{r}", font=f(True), fill=TOTAL, fmt=BAHT)

r += 2
b[f"A{r}"] = "สรุปตามกลุ่ม"
b[f"A{r}"].font = f(True)
r += 1
for j, h in enumerate(["กลุ่ม", "จำนวนคน", "งบต่อคน (บาท)", "งบประมาณ (บาท)",
                       "จ่ายจริงต่อคน (บาท)", "จ่ายจริง (บาท)", "ส่วนต่าง (บาท)"]):
    b.cell(row=r, column=j + 1, value=h)
style_range(b, f"A{r}:G{r}", font=f(True), fill=HEAD, align=Alignment(horizontal="center", wrap_text=True))
r += 1
g0 = r
for label, ref in [("นักศึกษา", f"={P}B3"), ("อาจารย์", f"={P}B4")]:
    b[f"A{r}"], b[f"B{r}"] = label, ref
    b[f"C{r}"] = f"=IF({PEOPLE}=0,0,$E${tot}/{PEOPLE})"
    b[f"D{r}"] = f"=B{r}*C{r}"
    b[f"E{r}"] = f"=IF({PEOPLE}=0,0,$H${tot}/{PEOPLE})"
    b[f"F{r}"] = f"=B{r}*E{r}"
    b[f"G{r}"] = f"=D{r}-F{r}"
    style_range(b, f"A{r}:G{r}", font=f(), fmt=BAHT)
    b[f"B{r}"].font = f(color=GREEN)
    r += 1
b[f"A{r}"] = "รวม"
for col in "BDFG":
    b[f"{col}{r}"] = f"=SUM({col}{g0}:{col}{r-1})"
style_range(b, f"A{r}:G{r}", font=f(True), fill=TOTAL, fmt=BAHT)

r += 2
for n in ["ส่วนต่างเป็นบวก = ใช้จ่ายน้อยกว่างบ (เหลือ) | ส่วนต่างเป็นลบ (ในวงเล็บ) = ใช้เกินงบ",
          "ช่องสีเหลือง = กรอกจำนวนคนที่กินจริง | ตัวเลขสีเขียว = ดึงมาจากชีต 'สมมติฐาน'",
          "ราคาจริงจากผู้ใช้: อาหารว่าง 40 บาท อาหารกลางวัน 55 บาท อาหารเย็น 55 บาท ต่อคนต่อมื้อ",
          "จ่ายจริงต่อคนของแต่ละกลุ่ม = ยอดจ่ายจริงรวม ÷ จำนวนคนทั้งหมด (เฉลี่ยเท่ากันทุกคน)"]:
    b[f"A{r}"] = n
    b[f"A{r}"].font = f(color="595959")
    r += 1

for col, w in zip("ABCDEFGHI", [16, 18, 14, 16, 16, 14, 15, 15, 16]):
    b.column_dimensions[col].width = w
b.freeze_panes = "A3"

# ---------------- Sheet 3: all possible amounts ----------------
c = wb.create_sheet("จำนวนเงินที่เป็นไปได้")
c["A1"] = "จำนวนเงินที่เป็นไปได้ทั้งหมด ตามจำนวนคนที่มากินจริง"
c["A1"].font = f(True, size=13)
c["B2"] = "ราคาตั้งงบ"
c.merge_cells("B2:G2")
c["I2"] = "ราคาจริง"
c.merge_cells("I2:L2")
heads = ["จำนวนคน", "อาหารว่าง 1 มื้อ", "อาหารกลางวัน 1 มื้อ", "อาหารเย็น 1 มื้อ",
         "ต่อวัน (ว่าง 2 + กลางวัน)", "วัน 24 ต.ค. (รวมอาหารเย็น)",
         "งบรวมทั้งโครงการ (บาท)", "ส่วนต่างจากงบเต็ม (บาท)",
         "อาหารว่าง 1 มื้อ", "อาหารกลางวัน 1 มื้อ", "อาหารเย็น 1 มื้อ",
         "จ่ายจริงทั้งโครงการ (บาท)", "ส่วนต่างจากงบเต็ม (บาท)"]
for j, h in enumerate(heads):
    c.cell(row=3, column=j + 1, value=h)
style_range(c, "A2:M2", font=f(True), fill=HEAD, align=Alignment(horizontal="center"))
style_range(c, "A3:M3", font=f(True), fill=HEAD,
            align=Alignment(horizontal="center", vertical="center", wrap_text=True))
c.row_dimensions[3].height = 32
first = 4
for i, n in enumerate(range(18, -1, -1)):
    r = first + i
    c[f"A{r}"] = n
    c[f"B{r}"] = f"=A{r}*{SNACK}"
    c[f"C{r}"] = f"=A{r}*{LUNCH}"
    c[f"D{r}"] = f"=A{r}*{DINNER}"
    c[f"E{r}"] = f"={P}$B$6*B{r}+{P}$B$8*C{r}"
    c[f"F{r}"] = f"=E{r}+D{r}"
    c[f"G{r}"] = f"=E{r}*{P}$B$11+D{r}"
    c[f"H{r}"] = f"=$G${first}-G{r}"
    c[f"I{r}"] = f"=A{r}*{SNACK_A}"
    c[f"J{r}"] = f"=A{r}*{LUNCH_A}"
    c[f"K{r}"] = f"=A{r}*{DINNER_A}"
    c[f"L{r}"] = f"=({P}$B$6*I{r}+{P}$B$8*J{r})*{P}$B$11+K{r}"
    c[f"M{r}"] = f"=$G${first}-L{r}"
    style_range(c, f"A{r}:M{r}", font=f(), fmt=BAHT)
    c[f"A{r}"].font = f(color=BLUE)
    if n == 18:
        style_range(c, f"A{r}:M{r}", font=f(True), fill=TOTAL, fmt=BAHT)
    c[f"A{r}"].number_format = "0"
    c[f"A{r}"].alignment = Alignment(horizontal="center")
lastc = first + 18
one = lastc - 1  # row for 1 person
r = lastc + 2
c[f"A{r}"] = "ส่วนต่างที่เกิดได้เมื่อขาดคน (ลดลงต่อครั้ง)"
c[f"A{r}"].font = f(True)
r += 1
c[f"A{r}"], c[f"G{r}"], c[f"L{r}"] = "กรณี", "ตามราคาตั้งงบ (บาท)", "ตามราคาจริง (บาท)"
style_range(c, f"A{r}:L{r}", font=f(True), fill=HEAD)
r += 1
cases = [
    ("ขาด 1 คน ในอาหารว่าง 1 มื้อ", f"={SNACK}", f"={SNACK_A}"),
    ("ขาด 1 คน ในอาหารกลางวัน 1 มื้อ", f"={LUNCH}", f"={LUNCH_A}"),
    ("ขาด 1 คน ในอาหารเย็น (24 ต.ค.)", f"={DINNER}", f"={DINNER_A}"),
    ("ขาด 1 คน ทั้งวัน (วันปกติ)", f"=E{one}", f"={P}$B$6*I{one}+{P}$B$8*J{one}"),
    ("ขาด 1 คน ทั้งวัน 24 ต.ค.", f"=F{one}", f"={P}$B$6*I{one}+{P}$B$8*J{one}+K{one}"),
    ("ขาด 1 คน ตลอดทั้งโครงการ", f"=G{one}", f"=L{one}"),
]
for lab, fx, fx_a in cases:
    c[f"A{r}"], c[f"G{r}"], c[f"L{r}"] = lab, fx, fx_a
    c.merge_cells(f"A{r}:F{r}")
    style_range(c, f"A{r}:L{r}", font=f(), fmt=BAHT)
    r += 1
r += 1
for n in [
    "ราคาตั้งงบ: ทุกราคาเป็นทวีคูณของ 50 บาท ยอดจึงเป็นได้ตั้งแต่ 0 ถึงงบเต็ม ทีละ 50 บาท",
    "ราคาจริง: ราคาเป็นทวีคูณของ 5 บาท (40, 55) ยอดจ่ายจริงจึงขยับทีละ 5 บาท",
    "ส่วนต่างจากงบเต็มในคอลัมน์ M = งบเต็ม (18 คน ราคาตั้งงบ) − จ่ายจริงตามจำนวนคน",
]:
    c[f"A{r}"] = "• " + n
    c[f"A{r}"].font = f(color="595959")
    r += 1
c[f"A{r}"] = "จำนวนยอดงบที่เป็นไปได้ตามราคาตั้งงบ (กรณี)"
c[f"G{r}"] = f"=G{first}/{SNACK}+1"
style_range(c, f"A{r}:G{r}", font=f(True), fmt=NUM, border=False)
c[f"G{r}"].comment = Comment("นับตั้งแต่ 0 ถึงงบเต็ม ทีละ 50 บาท", "budget")

for col, w in zip("ABCDEFGHIJKLM", [12, 14, 15, 14, 16, 18, 18, 18, 14, 15, 14, 18, 18]):
    c.column_dimensions[col].width = w
c.freeze_panes = "B4"

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
